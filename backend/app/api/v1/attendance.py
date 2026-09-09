"""Attendance, leave, and overtime (Phase 2) — org-scoped."""
from __future__ import annotations

import csv
import io
from datetime import date

from fastapi import APIRouter, Body, Depends, File, HTTPException, Request, Response, UploadFile
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.deps import require_org_access, require_permission
from app.models.attendance import AttendanceLog, LeaveType
from app.models.hr import LeaveRequest, OvertimeRequest
from app.models.person import Engagement
from app.models.user import User
from app.services import audit_service

router = APIRouter(prefix="/organizations/{organization_id}", tags=["attendance-leave"])

# Columns accepted by the attendance importer (also the template header order).
IMPORT_COLUMNS = ["employee_number", "log_date", "time_in", "time_out",
                  "hours_worked", "late_minutes", "overtime_hours", "status"]


def _to_num(v, default=0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def _ingest_rows(db: Session, organization_id: int, rows: list[dict]) -> dict:
    """Match rows to engagements by employee_number; upsert AttendanceLog by date."""
    emp_map = {
        e.employee_number: e.id
        for e in db.query(Engagement).filter(Engagement.organization_id == organization_id,
                                             Engagement.employee_number.isnot(None)).all()
    }
    imported, updated, errors = 0, 0, []
    for i, row in enumerate(rows, start=1):
        emp_no = str(row.get("employee_number", "")).strip()
        if not emp_no or emp_no not in emp_map:
            errors.append(f"row {i}: unknown employee_number '{emp_no}'")
            continue
        raw_date = str(row.get("log_date") or row.get("date") or "").strip()
        try:
            log_date = date.fromisoformat(raw_date)
        except ValueError:
            errors.append(f"row {i}: invalid date '{raw_date}' (expected YYYY-MM-DD)")
            continue
        eng_id = emp_map[emp_no]
        existing = (db.query(AttendanceLog)
                    .filter(AttendanceLog.engagement_id == eng_id, AttendanceLog.log_date == log_date).first())
        target = existing or AttendanceLog(organization_id=organization_id, engagement_id=eng_id, log_date=log_date)
        target.hours_worked = _to_num(row.get("hours_worked"), 8)
        target.late_minutes = int(_to_num(row.get("late_minutes"), 0))
        target.overtime_hours = _to_num(row.get("overtime_hours"), 0)
        target.status = str(row.get("status") or "PRESENT").strip() or "PRESENT"
        target.source = row.get("source", "CSV")
        if existing:
            updated += 1
        else:
            db.add(target)
            imported += 1
    db.commit()
    return {"imported": imported, "updated": updated, "errors": errors[:50], "error_count": len(errors)}


@router.get("/attendance/template", dependencies=[Depends(require_permission("attendance.view"))])
def attendance_template(organization_id: int, _: int = Depends(require_org_access)) -> Response:
    """Downloadable CSV template with the exact headers and a sample row."""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(IMPORT_COLUMNS)
    w.writerow(["EMP-1000", date.today().isoformat(), "08:00", "17:00", "8", "0", "0", "PRESENT"])
    return Response(content=buf.getvalue().encode("utf-8"), media_type="text/csv",
                    headers={"Content-Disposition": 'attachment; filename="attendance_template.csv"'})


@router.get("/attendance/format", dependencies=[Depends(require_permission("attendance.view"))])
def attendance_format(organization_id: int, _: int = Depends(require_org_access)) -> dict:
    return {
        "columns": IMPORT_COLUMNS,
        "notes": "One row per employee per day. Match is by employee_number. Dates are YYYY-MM-DD, "
                 "times HH:MM (optional). Re-importing the same employee+date updates that day.",
        "accepts": ["CSV (.csv)", "Excel (.xlsx)", "JSON via /attendance/device for biometric devices/APIs"],
    }


@router.post("/attendance/import", dependencies=[Depends(require_permission("attendance.edit"))])
async def import_attendance(organization_id: int, _: int = Depends(require_org_access),
                            file: UploadFile = File(...), db: Session = Depends(get_db)) -> dict:
    content = await file.read()
    name = (file.filename or "").lower()
    rows: list[dict] = []
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        header = None
        for r in ws.iter_rows(values_only=True):
            if header is None:
                header = [str(c).strip() if c is not None else "" for c in r]
                continue
            rows.append({header[i]: r[i] for i in range(min(len(header), len(r)))})
    else:
        text = content.decode("utf-8-sig", errors="replace")
        rows = list(csv.DictReader(io.StringIO(text)))
    return _ingest_rows(db, organization_id, rows)


@router.post("/attendance/device", dependencies=[Depends(require_permission("attendance.edit"))])
def import_attendance_device(organization_id: int, payload: dict = Body(...),
                             _: int = Depends(require_org_access), db: Session = Depends(get_db)) -> dict:
    """Generic ingestion endpoint for biometric devices / integrations.
    Body: {"logs": [{employee_number, log_date, time_in, time_out, hours_worked, ...}, ...]}"""
    logs = payload.get("logs") or []
    if not isinstance(logs, list):
        raise HTTPException(status_code=422, detail="'logs' must be a list")
    return _ingest_rows(db, organization_id, logs)


# --- Attendance --------------------------------------------------------------
@router.get("/attendance", dependencies=[Depends(require_permission("attendance.view"))])
def list_attendance(
    organization_id: int, engagement_id: int | None = None,
    _: int = Depends(require_org_access), db: Session = Depends(get_db),
):
    q = db.query(AttendanceLog).filter(AttendanceLog.organization_id == organization_id)
    if engagement_id:
        q = q.filter(AttendanceLog.engagement_id == engagement_id)
    rows = q.order_by(AttendanceLog.log_date.desc()).limit(200).all()
    return [
        {"id": r.id, "engagement_id": r.engagement_id, "log_date": r.log_date.isoformat(),
         "hours_worked": float(r.hours_worked or 0), "late_minutes": r.late_minutes,
         "overtime_hours": float(r.overtime_hours or 0), "source": r.source, "status": r.status}
        for r in rows
    ]


@router.post("/attendance", dependencies=[Depends(require_permission("attendance.edit"))])
def create_attendance(
    organization_id: int, payload: dict = Body(...),
    _: int = Depends(require_org_access), db: Session = Depends(get_db),
):
    log = AttendanceLog(
        organization_id=organization_id,
        engagement_id=payload["engagement_id"],
        log_date=date.fromisoformat(payload.get("log_date", date.today().isoformat())),
        hours_worked=payload.get("hours_worked", 8),
        late_minutes=payload.get("late_minutes", 0),
        overtime_hours=payload.get("overtime_hours", 0),
        source=payload.get("source", "MANUAL"),
        status=payload.get("status", "PRESENT"),
    )
    db.add(log)
    db.commit()
    return {"id": log.id}


# --- Leave -------------------------------------------------------------------
@router.get("/leave-types", dependencies=[Depends(require_permission("leave.view"))])
def leave_types(organization_id: int, _: int = Depends(require_org_access), db: Session = Depends(get_db)):
    rows = db.query(LeaveType).filter(LeaveType.organization_id == organization_id).all()
    return [{"id": r.id, "name": r.name, "default_credits": float(r.default_credits or 0), "paid": r.paid}
            for r in rows]


@router.get("/leave", dependencies=[Depends(require_permission("leave.view"))])
def list_leave(
    organization_id: int, status: str | None = None,
    _: int = Depends(require_org_access), db: Session = Depends(get_db),
):
    q = db.query(LeaveRequest).filter(LeaveRequest.organization_id == organization_id)
    if status:
        q = q.filter(LeaveRequest.status == status)
    rows = q.order_by(LeaveRequest.id.desc()).limit(200).all()
    return [
        {"id": r.id, "engagement_id": r.engagement_id, "leave_type": r.leave_type,
         "date_from": r.date_from.isoformat() if r.date_from else None,
         "date_to": r.date_to.isoformat() if r.date_to else None,
         "days": float(r.days or 0), "status": r.status}
        for r in rows
    ]


@router.post("/leave", dependencies=[Depends(require_permission("leave.apply"))])
def apply_leave(
    organization_id: int, payload: dict = Body(...),
    _: int = Depends(require_org_access), db: Session = Depends(get_db),
):
    lr = LeaveRequest(
        organization_id=organization_id, engagement_id=payload["engagement_id"],
        leave_type=payload.get("leave_type", "Vacation"),
        date_from=date.fromisoformat(payload["date_from"]) if payload.get("date_from") else None,
        date_to=date.fromisoformat(payload["date_to"]) if payload.get("date_to") else None,
        days=payload.get("days", 1), status="PENDING",
    )
    db.add(lr)
    db.commit()
    return {"id": lr.id, "status": lr.status}


@router.post("/leave/{leave_id}/decision", dependencies=[Depends(require_permission("leave.approve"))])
def decide_leave(
    organization_id: int, leave_id: int, request: Request, payload: dict = Body(...),
    _: int = Depends(require_org_access),
    user: User = Depends(require_permission("leave.approve")),
    db: Session = Depends(get_db),
):
    lr = db.get(LeaveRequest, leave_id)
    if not lr or lr.organization_id != organization_id:
        raise HTTPException(status_code=404, detail="Leave request not found")
    decision = payload.get("decision", "APPROVED").upper()
    if decision not in ("APPROVED", "REJECTED"):
        raise HTTPException(status_code=400, detail="decision must be APPROVED or REJECTED")
    lr.status = decision
    audit_service.record(db, action="leave.decision", user=user, organization_id=organization_id,
                         entity="leave_request", entity_id=leave_id, after={"status": decision},
                         ip_address=request.client.host if request.client else None, commit=False)
    db.commit()
    return {"id": lr.id, "status": lr.status}


# --- Overtime ----------------------------------------------------------------
@router.get("/overtime", dependencies=[Depends(require_permission("attendance.view"))])
def list_overtime(
    organization_id: int, status: str | None = None,
    _: int = Depends(require_org_access), db: Session = Depends(get_db),
):
    q = db.query(OvertimeRequest).filter(OvertimeRequest.organization_id == organization_id)
    if status:
        q = q.filter(OvertimeRequest.status == status)
    rows = q.order_by(OvertimeRequest.id.desc()).limit(200).all()
    return [{"id": r.id, "engagement_id": r.engagement_id,
             "ot_date": r.ot_date.isoformat() if r.ot_date else None,
             "hours": float(r.hours or 0), "status": r.status} for r in rows]


@router.post("/overtime/{ot_id}/decision", dependencies=[Depends(require_permission("attendance.approve"))])
def decide_overtime(
    organization_id: int, ot_id: int, payload: dict = Body(...),
    _: int = Depends(require_org_access), db: Session = Depends(get_db),
):
    ot = db.get(OvertimeRequest, ot_id)
    if not ot or ot.organization_id != organization_id:
        raise HTTPException(status_code=404, detail="Overtime request not found")
    ot.status = payload.get("decision", "APPROVED").upper()
    db.commit()
    return {"id": ot.id, "status": ot.status}
